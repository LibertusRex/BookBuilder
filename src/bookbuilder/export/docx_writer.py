"""Exportação para DOCX (Word).

Princípios:

* Texto real e editável, em ordem de leitura, com estilos semânticos
  (Normal, Título 1–6, Citação, Legenda, Parágrafo da Lista, Texto de nota).
* Cabeçalhos e rodapés vão para as partes nativas de cabeçalho/rodapé do
  Word (``header*.xml``/``footer*.xml``), nunca para o corpo. Números de
  página viram o campo PAGE. Páginas pares/ímpares e primeira página
  diferente são usadas quando o original alterna.
* Notas de rodapé viram notas nativas do Word (``footnotes.xml``), ligadas à
  chamada no texto.
* Nada de caixas de texto flutuantes nem páginas como imagem; quadros com
  moldura viram parágrafos com borda.
"""

from __future__ import annotations

import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from ..document import flow_pages
from ..layout.cleanup import TOC
from ..layout.noise import COVER
from ..layout.structure import FLOW_ROLES, box_group
from ..layout.running import normalize
from ..model import Block, Page, Project, Role, Run, Style
from ..typography import fonts

TOC_STYLE = "Sumário"

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"


@dataclass
class DocxResult:
    path: str
    paragraphs: int = 0
    footnotes: int = 0
    sections: int = 0
    images: int = 0
    tables: int = 0
    notes: list[str] = field(default_factory=list)


class _Geometry:
    """Conversões de pixels da imagem para pontos tipográficos."""

    def __init__(self, project: Project, pages: list[Page]):
        self.ppp = float(project.analysis.get("px_per_pt") or 300 / 72)
        widths = [p.width for p in pages if p.width] or [2480]
        heights = [p.height for p in pages if p.height] or [3508]
        self.page_w = statistics.median(widths) / self.ppp
        self.page_h = statistics.median(heights) / self.ppp
        flow = [b for p in pages for b in p.blocks if b.role in FLOW_ROLES and len(b.lines) >= 2]
        lefts = [b.bbox[0] for b in flow] or [widths[0] * 0.1]
        rights = [b.bbox[2] for b in flow] or [widths[0] * 0.9]
        tops = [min(b.bbox[1] for b in p.blocks if b.role in FLOW_ROLES or b.role == Role.TITLE)
                for p in pages if any(b.role in FLOW_ROLES or b.role == Role.TITLE for b in p.blocks)] or [heights[0] * 0.1]
        bottoms = [max(b.bbox[3] for b in p.blocks if b.role in FLOW_ROLES or b.role == Role.FOOTNOTE)
                   for p in pages if any(b.role in FLOW_ROLES for b in p.blocks)] or [heights[0] * 0.9]
        self.margin_l = max(36.0, statistics.median(lefts) / self.ppp)
        self.margin_r = max(36.0, self.page_w - statistics.median(rights) / self.ppp)
        self.margin_t = max(36.0, statistics.median(tops) / self.ppp)
        self.margin_b = max(36.0, self.page_h - statistics.median(bottoms) / self.ppp)
        if self.page_w - self.margin_l - self.margin_r < 144:
            self.margin_l = self.margin_r = 54.0
        heads = [b.bbox[1] for p in pages for b in p.blocks if b.role == Role.HEADER]
        foots = [b.bbox[3] for p in pages for b in p.blocks if b.role in (Role.FOOTER, Role.PAGE_NUMBER)]
        self.header_dist = max(12.0, statistics.median(heads) / self.ppp) if heads else 36.0
        self.footer_dist = max(12.0, self.page_h - statistics.median(foots) / self.ppp) if foots else 36.0
        self.header_dist = min(self.header_dist, self.margin_t - 6)
        self.footer_dist = min(self.footer_dist, self.margin_b - 6)
        gaps = [b.space_before for p in pages for b in p.blocks if b.role == Role.BODY and b.space_before > 0]
        self.normal_gap = statistics.median(gaps) if gaps else 0.0

    def pt(self, px: float) -> float:
        return px / self.ppp

    @property
    def text_width(self) -> float:
        return self.page_w - self.margin_l - self.margin_r


# ---------------------------------------------------------------------------
# Estilos


def _style_font(style, st: Style | None, size: float | None = None) -> None:
    fam = _usable_family(st) if st else "Georgia"
    style.font.name = fam
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), fam)
    style.font.size = Pt(size or (st.size if st else 11))
    if st is not None:
        style.font.bold = st.bold
        style.font.italic = st.italic


def _usable_family(st: Style) -> str:
    fam = st.family
    if fonts.installed_faces() and not fonts.is_installed(fam):
        return st.detected_family or {"serif": "Georgia", "sans": "Segoe UI", "mono": "Consolas"}[st.category]
    return fam


def _get_or_add_style(doc, name: str, kind=WD_STYLE_TYPE.PARAGRAPH, base: str | None = "Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        s = doc.styles.add_style(name, kind)
        if base and kind == WD_STYLE_TYPE.PARAGRAPH:
            s.base_style = doc.styles[base]
        return s


def _dominant(blocks: list[Block], styles: dict[str, Style]) -> Style | None:
    c = Counter(b.style_id for b in blocks if b.style_id in styles)
    return styles[c.most_common(1)[0][0]] if c else None


def _setup_styles(doc, project: Project, pages: list[Page], geo: _Geometry) -> dict[str, str]:
    """Cria/ajusta os estilos do Word. Retorna estilo detectado → nome do
    estilo de parágrafo no Word."""
    styles = project.styles
    all_blocks = [b for p in pages for b in p.blocks]
    body = next((s for s in styles.values() if s.name == "Corpo do texto"), None)
    normal = doc.styles["Normal"]
    _style_font(normal, body)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pitches = [b.line_pitch for b in all_blocks if b.role == Role.BODY and b.line_pitch]
    if pitches and body:
        line_pt = statistics.median(pitches) / geo.ppp
        pf.line_spacing = max(1.0, min(2.5, line_pt / (body.size * 1.17)))
    else:
        pf.line_spacing = 1.15
    pf.widow_control = True
    mapping: dict[str, str] = {}
    if body:
        mapping[body.id] = "Normal"
    # Títulos: estilo dominante de cada nível.
    for lvl in range(1, 7):
        bl = [b for b in all_blocks if b.role == Role.TITLE and b.heading_level == lvl]
        if not bl:
            continue
        hs = doc.styles[f"Heading {lvl}"]
        st = _dominant(bl, styles)
        _style_font(hs, st)
        hs.font.bold = True  # títulos, subtítulos e seções sempre em negrito
        hs.font.color.rgb = None
        hpf = hs.paragraph_format
        hpf.space_before = Pt(max(6.0, (st.size if st else 12) * 1.2))
        hpf.space_after = Pt(max(3.0, (st.size if st else 12) * 0.5))
        hpf.keep_with_next = True
        al = Counter(b.align for b in bl).most_common(1)[0][0]
        hpf.alignment = ALIGN.get(al, WD_ALIGN_PARAGRAPH.LEFT)
    # Demais estilos detectados: estilos próprios baseados em Normal.
    for st in styles.values():
        if st.id in mapping or st.role == Role.TITLE:
            continue
        base = {Role.FOOTNOTE: None, Role.CAPTION: "Caption"}.get(st.role, "Normal")
        if st.role in (Role.HEADER, Role.FOOTER, Role.PAGE_NUMBER):
            name = "Header" if st.role == Role.HEADER else "Footer"
        elif st.role == Role.FOOTNOTE:
            name = "Footnote Text"
        else:
            name = st.name
        s = _get_or_add_style(doc, name, base=base if base else "Normal")
        _style_font(s, st)
        mapping[st.id] = name
    # Quadros: borda simples em volta dos parágrafos consecutivos.
    q = _get_or_add_style(doc, "Quadro")
    qst = _dominant([b for b in all_blocks if b.box_id], styles)
    _style_font(q, qst)
    _add_borders(q.element.get_or_add_pPr())
    q.paragraph_format.left_indent = Pt(6)
    q.paragraph_format.right_indent = Pt(6)
    _get_or_add_style(doc, "Quote")
    _get_or_add_style(doc, "Caption")
    # Sumário do livro: todas as entradas no mesmo formato, número da página
    # na margem direita com pontilhado.
    toc = _get_or_add_style(doc, TOC_STYLE)
    _style_font(toc, body)
    toc.font.bold = False
    toc.font.italic = False
    tpf = toc.paragraph_format
    tpf.first_line_indent = Pt(0)
    tpf.left_indent = Pt(0)
    tpf.space_before = Pt(0)
    tpf.space_after = Pt(max(2.0, (body.size if body else 11) * 0.3))
    tpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tpf.tab_stops.add_tab_stop(Pt(geo.text_width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    fn = _get_or_add_style(doc, "Footnote Text")
    fst = _dominant([b for b in all_blocks if b.role == Role.FOOTNOTE], styles)
    if fst:
        _style_font(fn, fst)
    else:
        fn.font.size = Pt(max(7.0, (body.size if body else 11) * 0.8))
    # Cabeçalho e rodapé: tabulações central e à direita na largura do texto
    # desta página (as do Word são da folha Carta).
    for name in ("Header", "Footer"):
        # Letra bem pequena, na fonte do corpo (o tamanho medido na foto
        # fica grande demais na página do Word).
        hf = doc.styles[name]
        hf.font.size = Pt(max(5.0, round((body.size if body else 11.0) * 0.45 * 2) / 2))
        hf.font.bold = False
        hf.font.name = None
        rpr = hf.element.get_or_add_rPr()
        for rf in rpr.findall(qn("w:rFonts")):
            rpr.remove(rf)
        ts = hf.paragraph_format.tab_stops
        ts.clear_all()
        ts.add_tab_stop(Pt(geo.text_width / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Pt(geo.text_width), WD_TAB_ALIGNMENT.RIGHT)
    fr = _get_or_add_style(doc, "Footnote Reference", kind=WD_STYLE_TYPE.CHARACTER, base=None)
    fr.font.superscript = True
    return mapping


def _add_borders(ppr) -> None:
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), "808080")
        bdr.append(el)
    # Ordem exigida pelo esquema do Word: pBdr antes de shd, tabs, spacing, ind, jc...
    ppr.insert_element_before(bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                              "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                              "w:pPrChange")


# ---------------------------------------------------------------------------
# Runs


class _Notes:
    def __init__(self):
        self.items: list[tuple[int, list[Run]]] = []
        self.by_key: dict[str, int] = {}
        self.next_id = 1

    def id_for(self, key: str) -> int:
        if key not in self.by_key:
            self.by_key[key] = self.next_id
            self.next_id += 1
        return self.by_key[key]


def _add_runs(par, runs: list[Run], notes: _Notes, note_texts: dict[str, list[Run]], highlight: bool) -> None:
    for r in runs:
        if r.note_ref and r.note_ref in note_texts:
            nid = notes.id_for(r.note_ref)
            _footnote_reference(par, nid)
            continue
        parts = re.split(r"(\t|\n)", r.text)
        for part in parts:
            if not part:
                continue
            if part == "\t":
                par.add_run().add_tab()
                continue
            if part == "\n":
                par.add_run().add_break(WD_BREAK.LINE)
                continue
            if r.link:
                _hyperlink(par, part, r, r.link)
                continue
            pieces = URL_RE.split(part)
            if len(pieces) > 1:
                for k, piece in enumerate(pieces):
                    if not piece:
                        continue
                    if k % 2 == 1:
                        _hyperlink(par, piece, r)
                    else:
                        _plain_run(par, piece, r, highlight)
                continue
            _plain_run(par, part, r, highlight)


# Endereço web no texto (o ponto ou parêntese final não faz parte dele).
URL_RE = re.compile(r"((?:https?://|www\.)[^\s<>“”\"]+?)(?=[.,;:)\]]*(?:\s|$))")


def url_href(u: str) -> str:
    return u if re.match(r"https?://", u) else "http://" + u


def _plain_run(par, part: str, r: Run, highlight: bool) -> None:
    run = par.add_run(part)
    f = run.font
    if r.bold:
        f.bold = True
    if r.italic:
        f.italic = True
    if r.underline:
        f.underline = True
    if r.superscript:
        f.superscript = True
    if r.subscript:
        f.subscript = True
    if r.low_conf and highlight:
        f.highlight_color = WD_COLOR_INDEX.YELLOW


def _hyperlink(par, url: str, r: Run, target: str | None = None) -> None:
    """Hiperlink externo (azul, sublinhado); ``url`` é o texto mostrado e
    ``target`` o destino, quando diferente dele."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    rid = par.part.relate_to(target or url_href(url), RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    if r.bold:
        rpr.append(OxmlElement("w:b"))
    if r.italic:
        rpr.append(OxmlElement("w:i"))
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    par._p.append(h)


def _footnote_reference(par, nid: int) -> None:
    run = par.add_run()
    rpr = run._r.get_or_add_rPr()
    rs = OxmlElement("w:rStyle")
    rs.set(qn("w:val"), "FootnoteReference")
    rpr.append(rs)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(nid))
    run._r.append(ref)


def _xml_escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _footnotes_xml(notes: _Notes, note_texts: dict[str, list[Run]], highlight: bool) -> bytes:
    parts = [
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W_NS}">',
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:separator/></w:r></w:p></w:footnote>',
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
    ]
    for key, nid in sorted(notes.by_key.items(), key=lambda kv: kv[1]):
        runs = note_texts.get(key, [])
        body = []
        for r in runs:
            if r.note_ref:
                continue
            props = []
            if r.bold:
                props.append("<w:b/>")
            if r.italic:
                props.append("<w:i/>")
            if r.underline:
                props.append('<w:u w:val="single"/>')
            if r.superscript:
                props.append('<w:vertAlign w:val="superscript"/>')
            if r.subscript:
                props.append('<w:vertAlign w:val="subscript"/>')
            if r.low_conf and highlight:
                props.append('<w:highlight w:val="yellow"/>')
            rpr = f"<w:rPr>{''.join(props)}</w:rPr>" if props else ""
            for part in re.split(r"(\t|\n)", r.text):
                if not part:
                    continue
                if part == "\t":
                    body.append(f"<w:r>{rpr}<w:tab/></w:r>")
                elif part == "\n":
                    body.append(f"<w:r>{rpr}<w:br/></w:r>")
                else:
                    body.append(f'<w:r>{rpr}<w:t xml:space="preserve">{_xml_escape(part)}</w:t></w:r>')
        parts.append(
            f'<w:footnote w:id="{nid}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
            '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
            '<w:r><w:t xml:space="preserve"> </w:t></w:r>' + "".join(body) + "</w:p></w:footnote>"
        )
    parts.append("</w:footnotes>")
    return "".join(parts).encode("utf-8")


def _attach_footnotes_part(doc, xml: bytes) -> None:
    main = doc.part
    for rel in main.rels.values():
        if rel.reltype == RT.FOOTNOTES:
            rel.target_part._blob = xml  # type: ignore[attr-defined]
            return
    part = Part(PackURI("/word/footnotes.xml"), FOOTNOTES_CT, xml, main.package)
    main.relate_to(part, RT.FOOTNOTES)


# ---------------------------------------------------------------------------
# Cabeçalhos e rodapés


@dataclass
class _Margin:
    """Conteúdo de cabeçalho ou rodapé de uma página: parágrafos (texto com
    "#" no lugar do número de página) e alinhamentos."""

    items: list[tuple[str, str, bool]]  # (texto, alinhamento, contém número de página)

    def signature(self) -> str:
        return "|".join(normalize(t) + a for t, a, _ in self.items)


def _margin_of(p: Page, blocks: list[Block], notes: list[str]) -> _Margin:
    items = []
    # Linha a linha (blocos que se sobrepõem na vertical), da esquerda para a direita.
    rows: list[list[Block]] = []
    for b in sorted(blocks, key=lambda b: (b.bbox[1], b.bbox[0])):
        row = rows[-1] if rows else None
        if row and min(b.bbox[3], max(x.bbox[3] for x in row)) - max(b.bbox[1], min(x.bbox[1] for x in row)) > 0:
            row.append(b)
        else:
            rows.append([b])
    for b in [b for row in rows for b in sorted(row, key=lambda b: b.bbox[0])]:
        if b.surya_html and b.runs:
            text = " ".join("".join(r.text for r in b.runs).split())  # leitura fundida Tesseract × Surya
        else:
            text = " ".join(ln.text for ln in b.lines).strip()
        if not text:
            continue
        cx = (b.bbox[0] + b.bbox[2]) / 2 / max(1, p.width)
        align = "center" if 0.38 < cx < 0.62 else ("right" if cx >= 0.62 else "left")
        if b.role == Role.PAGE_NUMBER:
            text = re.sub(r"\d+|\b[ivxlcdm]+\b$", "#", text, count=1, flags=re.I)
            # "Page 168 of 2245" / "página 3 de 120": o total vira o campo NUMPAGES (o do Word).
            text = re.sub(r"(#\s*(?:of|de|/)\s*)\d+", "\\1@", text, count=1, flags=re.I)
            items.append((text, align, True))
        else:
            if re.search(r"\d", text) and re.fullmatch(r"\d{1,3}\s*%", text):
                if "progresso" not in notes:
                    notes.append("progresso")
                continue  # progresso de leitura do aplicativo: muda a cada página
            items.append((text, align, False))
    return _Margin(items)


def _fill_margin(container, margin: _Margin | None) -> None:
    for p in list(container.paragraphs):
        p._p.getparent().remove(p._p)
    style = "Header" if "Header" in type(container).__name__ else "Footer"

    def new_par(text: str = ""):
        par = container.add_paragraph(text)
        par.style = style
        return par

    if margin is None or not margin.items:
        new_par("")
        return
    by_align: dict[str, tuple[str, bool]] = {}
    for t, a, has in margin.items:
        if a in by_align:
            prev_t, prev_has = by_align[a]
            by_align[a] = (prev_t + " " + t, prev_has or has)
        else:
            by_align[a] = (t, has)
    if len(by_align) == 1:
        a, (t, has) = next(iter(by_align.items()))
        par = new_par()
        par.alignment = ALIGN[a]
        _margin_text(par, t, has)
        return
    # Vários itens na mesma faixa (ex.: título à esquerda, número à direita):
    # um parágrafo com tabulações central e à direita.
    par = new_par()
    _tab_stops(par)
    if "left" in by_align:
        _margin_text(par, *by_align["left"])
    if "center" in by_align:
        par.add_run().add_tab()
        _margin_text(par, *by_align["center"])
    if "right" in by_align:
        par.add_run().add_tab()
        if "center" not in by_align:
            par.add_run().add_tab()
        _margin_text(par, *by_align["right"])


def _tab_stops(par) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT

    section_width = None
    try:
        sec = par.part.package.main_document_part.document.sections[0]  # parágrafo de cabeçalho/rodapé
        section_width = sec.page_width - sec.left_margin - sec.right_margin
    except Exception:  # noqa: BLE001
        pass
    if section_width:
        ts = par.paragraph_format.tab_stops
        ts.add_tab_stop(Emu(int(section_width / 2)), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Emu(int(section_width)), WD_TAB_ALIGNMENT.RIGHT)


def _margin_text(par, text: str, has_number: bool) -> None:
    if not has_number:
        par.add_run(text)
        return
    for k, piece in enumerate(re.split(r"([#@])", text)):
        if piece == "#":
            _page_field(par)
        elif piece == "@":
            _page_field(par, "NUMPAGES")
        elif piece:
            par.add_run(piece)


def _page_field(par, field: str = "PAGE") -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {field} ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


# ---------------------------------------------------------------------------


def write(project: Project, store, out_path: str) -> DocxResult:
    pages = flow_pages(project)
    res = DocxResult(path=out_path)
    if not pages:
        raise ValueError("Nenhuma página processada para exportar.")
    opts = project.export
    geo = _Geometry(project, pages)
    doc = Document()
    _core_properties(doc, project)
    mapping = _setup_styles(doc, project, pages, geo)

    sec = doc.sections[0]
    _page_setup(sec, geo)

    # Conteúdo de margem por página.
    margin_notes: list[str] = []
    heads: list[_Margin] = []
    foots: list[_Margin] = []
    for p in pages:
        hb = [b for b in p.blocks if b.role in (Role.HEADER, Role.FOOTER, Role.PAGE_NUMBER) and (b.bbox[1] + b.bbox[3]) / 2 < p.height / 2]
        fb = [b for b in p.blocks if b.role in (Role.HEADER, Role.FOOTER, Role.PAGE_NUMBER) and (b.bbox[1] + b.bbox[3]) / 2 >= p.height / 2]
        heads.append(_margin_of(p, hb, margin_notes))
        foots.append(_margin_of(p, fb, margin_notes))
    if "progresso" in margin_notes:
        res.notes.append("O rodapé com a porcentagem de leitura do aplicativo (ex.: “6%”) muda a cada página e não tem "
                         "equivalente no Word; foi omitido do rodapé.")

    # Notas de rodapé: texto de cada nota (continuações unidas).
    note_texts: dict[str, list[Run]] = {}
    for p in pages:
        for b in p.ordered_blocks():
            if b.role == Role.FOOTNOTE and b.note_label:
                if b.note_label in note_texts and b.continues_prev:
                    note_texts[b.note_label] += [Run(" ")] + b.runs
                else:
                    note_texts[b.note_label] = list(b.runs)
    notes = _Notes()
    highlight = opts.highlight_low_conf

    body_container = doc
    cur_par = None
    sections_sig: str | None = None
    section_start_pages: list[tuple[object, int]] = [(sec, 0)]
    figures_dir = Path(store.figures_dir) if store is not None else None
    after_cover = False
    breaks = _scan_breaks(pages)
    for pi, p in enumerate(pages):
        sig = heads[pi].signature() + "||" + foots[pi].signature()
        if opts.docx_page_breaks and pi > 0:
            if sig != sections_sig and opts.docx_headers_footers:
                new = doc.add_section(WD_SECTION.NEW_PAGE)
                _page_setup(new, geo)
                section_start_pages.append((new, pi))
                cur_par_page_break = False
            else:
                cur_par_page_break = True
        else:
            cur_par_page_break = False
        sections_sig = sig
        first_on_page = True
        last_flow_par = None
        for b in p.ordered_blocks():
            if b.role.is_running or b.role == Role.FOOTNOTE:
                continue
            group = box_group(b)
            if b.continues_prev and cur_par is not None and b.role in FLOW_ROLES:
                if cur_par_page_break and first_on_page:
                    cur_par.add_run().add_break(WD_BREAK.PAGE)
                elif "juntar_sem_espaco" not in b.flags:
                    cur_par.add_run(" ")
                _add_runs(cur_par, b.runs, notes, note_texts, highlight)
                if b.align == "justify":
                    cur_par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                first_on_page = False
                last_flow_par = cur_par
                continue
            if b.role == Role.FIGURE:
                par = _add_figure(body_container, b, figures_dir, geo, res)
            elif b.role == Role.TABLE:
                par = _add_table(doc, b, figures_dir, geo, res, mapping)
            else:
                par = body_container.add_paragraph()
                _format_paragraph(par, b, mapping, geo, group, project)
                _add_runs(par, b.runs, notes, note_texts, highlight)
                _direct_font(par, b, project, mapping)
                res.paragraphs += 1
            if par is not None and ((first_on_page and cur_par_page_break) or after_cover or b.id in breaks):
                par.paragraph_format.page_break_before = True
                after_cover = False
            if COVER in b.flags and par is not None:
                after_cover = True  # a capa fica sozinha na primeira página
            first_on_page = False
            cur_par = par if b.role in FLOW_ROLES else None
            if b.role in FLOW_ROLES:
                last_flow_par = par
        # Notas da página sem chamada encontrada: referência no fim do último parágrafo.
        for b in p.ordered_blocks():
            if b.role == Role.FOOTNOTE and b.note_label and not b.continues_prev and b.note_label not in notes.by_key:
                target = last_flow_par or cur_par
                if target is None:
                    target = body_container.add_paragraph()
                nid = notes.id_for(b.note_label)
                _footnote_reference(target, nid)
                res.notes.append(f"Nota de rodapé sem chamada localizada em {p.source.display_name}: "
                                 "a referência foi colocada no fim do último parágrafo da página.")
    # Cabeçalhos/rodapés nas partes nativas.
    if opts.docx_headers_footers:
        if opts.docx_page_breaks:
            starts = section_start_pages + [(None, len(pages))]
            for (s, a), (_, bnd) in zip(starts, starts[1:]):
                _apply_margins(doc, s, heads[a:bnd], foots[a:bnd], pages[a:bnd], res)
        else:
            _apply_margins(doc, doc.sections[0], heads, foots, pages, res, continuous=True)
    res.sections = len(doc.sections)

    if notes.by_key:
        _attach_footnotes_part(doc, _footnotes_xml(notes, note_texts, highlight))
        res.footnotes = len(notes.by_key)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return res


def _scan_breaks(pages: list[Page]) -> set[str]:
    """Livro digitalizado: o sumário e cada capítulo começam em página nova
    (o título do sumário, o primeiro título depois dele e os títulos de
    nível 1). Devolve os ids dos blocos que levam a quebra antes."""
    out: set[str] = set()
    seq = [b for p in pages if p.light for b in p.ordered_blocks()
           if not b.role.is_running and b.role not in (Role.FOOTNOTE, Role.ARTIFACT)]
    if not seq:
        return out
    toc = [k for k, b in enumerate(seq) if TOC in b.flags]
    if toc:
        k0 = toc[0]
        if k0 > 0 and seq[k0 - 1].role == Role.TITLE:
            out.add(seq[k0 - 1].id)
        nxt = next((b for b in seq[toc[-1] + 1:] if b.role == Role.TITLE), None)
        if nxt is not None:
            out.add(nxt.id)
    out |= {b.id for b in seq[1:] if b.role == Role.TITLE and b.heading_level == 1}
    return out


def _core_properties(doc, project: Project) -> None:
    cp = doc.core_properties
    cp.title = project.export.title or project.export.base_name
    cp.author = project.export.author or ""
    cp.language = project.export.language or "pt-BR"
    cp.comments = "Gerado pelo BookBuilder a partir de imagens de páginas (OCR local)."
    # Idioma padrão do texto, para revisão ortográfica.
    lang = project.export.language or "pt-BR"
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    el = OxmlElement("w:lang")
    el.set(qn("w:val"), lang)
    rpr.append(el)


def _page_setup(sec, geo: _Geometry) -> None:
    sec.page_width = Pt(geo.page_w)
    sec.page_height = Pt(geo.page_h)
    sec.left_margin = Pt(geo.margin_l)
    sec.right_margin = Pt(geo.margin_r)
    sec.top_margin = Pt(geo.margin_t)
    sec.bottom_margin = Pt(geo.margin_b)
    sec.header_distance = Pt(max(12.0, geo.header_dist))
    sec.footer_distance = Pt(max(12.0, geo.footer_dist))


def _format_paragraph(par, b: Block, mapping: dict[str, str], geo: _Geometry, group: str | None, project: Project) -> None:
    doc_styles = par.part.document.styles
    if TOC in b.flags and any("\t" in r.text for r in b.runs):
        par.style = doc_styles[TOC_STYLE]
        extra = b.space_before - geo.normal_gap
        if extra > geo.ppp * 8:  # separação entre grupos de entradas
            par.paragraph_format.space_before = Pt(round(min(18.0, geo.pt(extra)), 1))
        return
    if b.role == Role.TITLE and b.heading_level:
        par.style = doc_styles[f"Heading {min(6, b.heading_level)}"]
        if group:
            # Fica dentro da moldura do quadro (mesma borda e recuos).
            _add_borders(par._p.get_or_add_pPr())
            par.paragraph_format.left_indent = Pt(6)
            par.paragraph_format.right_indent = Pt(6)
    elif group:
        par.style = doc_styles["Quadro"]
    elif b.role == Role.CAPTION:
        par.style = doc_styles["Caption"]
    elif b.role == Role.QUOTE:
        par.style = doc_styles["Quote"]
    elif b.role == Role.LIST_ITEM:
        try:
            par.style = doc_styles["List Paragraph"]
        except KeyError:
            pass
    elif b.style_id and b.style_id in mapping:
        try:
            par.style = doc_styles[mapping[b.style_id]]
        except KeyError:
            pass
    pf = par.paragraph_format
    pf.alignment = ALIGN.get(b.align, WD_ALIGN_PARAGRAPH.LEFT)
    if b.role in FLOW_ROLES or b.role == Role.CAPTION:
        if abs(b.indent_first) > 2:
            pf.first_line_indent = Pt(round(geo.pt(b.indent_first), 1))
        if b.indent_left > 6 and not group:
            pf.left_indent = Pt(round(geo.pt(b.indent_left), 1))
        if b.indent_right > geo.ppp * 12 and not group:
            pf.right_indent = Pt(round(geo.pt(b.indent_right), 1))
        extra = b.space_before - geo.normal_gap
        if extra > geo.ppp * 4:
            pf.space_before = Pt(round(min(36.0, geo.pt(extra)), 1))


def _direct_font(par, b: Block, project: Project, mapping: dict[str, str]) -> None:
    """Um bloco cujo estilo detectado não é o estilo do Word que ele usa
    (ex.: um título de nível 2 numa fonte diferente dos demais) recebe a
    fonte e o tamanho do seu estilo como formatação direta."""
    st = project.styles.get(b.style_id or "")
    if st is None or (par.style is not None and par.style.name == TOC_STYLE):
        return
    word_style = par.style.name if par.style is not None else ""
    if mapping.get(st.id) == word_style:
        return
    fam = _usable_family(st)
    for r in par.runs:
        r.font.name = fam
        rpr = r._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(attr), fam)
        if not r.font.superscript and not r.font.subscript:
            r.font.size = Pt(st.size)


def _add_figure(container, b: Block, figures_dir: Path | None, geo: _Geometry, res: DocxResult):
    if not b.image_file or figures_dir is None:
        return None
    path = figures_dir / b.image_file
    if not path.is_file():
        return None
    par = container.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = b.bbox[2] - b.bbox[0], b.bbox[3] - b.bbox[1]
    width_pt = min(geo.text_width, geo.pt(w))
    max_h = geo.page_h - geo.margin_t - geo.margin_b - 12  # cabe na página (capa)
    if h > 0 and width_pt * h / w > max_h:
        width_pt = max_h * w / h
    par.add_run().add_picture(str(path), width=Pt(width_pt))
    res.images += 1
    return par


def _add_table(doc, b: Block, figures_dir: Path | None, geo: _Geometry, res: DocxResult, mapping):
    if not b.table:
        # Tabela não reconhecida com confiança: imagem + texto do OCR.
        par = _add_figure(doc, b, figures_dir, geo, res)
        txt = doc.add_paragraph(" / ".join(ln.text for ln in b.lines))
        for r in txt.runs:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        res.notes.append("Uma tabela não foi reconhecida com segurança: foi inserida como imagem, seguida do texto lido.")
        return par
    rows = max(c.row for c in b.table) + 1
    cols = max(c.col for c in b.table) + 1
    t = doc.add_table(rows=rows, cols=cols)
    t.style = doc.styles["Table Grid"]
    for c in b.table:
        cell = t.cell(c.row, c.col)
        cell.text = c.text
    res.tables += 1
    return None


def _apply_margins(doc, sec, heads: list[_Margin], foots: list[_Margin], pages: list[Page], res: DocxResult,
                   continuous: bool = False) -> None:
    """Configura cabeçalho/rodapé de uma seção, com primeira página e
    pares/ímpares diferentes quando o original alterna."""
    if sec is None or not pages:
        return
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False

    def common(ms: list[_Margin]) -> _Margin | None:
        sigs = Counter(m.signature() for m in ms if m.items)
        if not sigs:
            return None
        top = sigs.most_common(1)[0][0]
        return next(m for m in ms if m.signature() == top)

    def parity_split(ms: list[_Margin]) -> tuple[_Margin | None, _Margin | None]:
        if continuous:
            nums = [p.printed_number for p in pages]
            if all(n is not None and n.isdigit() for n in nums):
                odd = [m for m, n in zip(ms, nums) if int(n) % 2 == 1]
                even = [m for m, n in zip(ms, nums) if int(n) % 2 == 0]
            else:
                odd, even = ms[0::2], ms[1::2]
        else:
            odd, even = ms[0::2], ms[1::2]
        return common(odd), common(even)

    h_odd, h_even = parity_split(heads)
    f_odd, f_even = parity_split(foots)
    alternating = any(
        a is not None and b is not None and a.signature() != b.signature()
        for a, b in ((h_odd, h_even), (f_odd, f_even))
    )
    first_h = heads[0] if heads else None
    first_f = foots[0] if foots else None
    h_main, f_main = common(heads), common(foots)
    first_differs = (first_h is not None and h_main is not None and first_h.signature() != h_main.signature()) or (
        first_f is not None and f_main is not None and first_f.signature() != f_main.signature())

    if alternating:
        doc.settings.odd_and_even_pages_header_footer = True
        _fill_margin(sec.header, h_odd or h_main)
        _fill_margin(sec.footer, f_odd or f_main)
        sec.even_page_header.is_linked_to_previous = False
        sec.even_page_footer.is_linked_to_previous = False
        _fill_margin(sec.even_page_header, h_even or h_main)
        _fill_margin(sec.even_page_footer, f_even or f_main)
        res.notes.append("Cabeçalhos/rodapés diferentes em páginas pares e ímpares.")
    else:
        _fill_margin(sec.header, h_main)
        _fill_margin(sec.footer, f_main)
    cover = any(COVER in b.flags for b in pages[0].blocks)  # capa: sem cabeçalho nem rodapé
    if (first_differs or cover) and len(pages) > 1:
        sec.different_first_page_header_footer = True
        sec.first_page_header.is_linked_to_previous = False
        sec.first_page_footer.is_linked_to_previous = False
        _fill_margin(sec.first_page_header, None if cover else first_h)
        _fill_margin(sec.first_page_footer, None if cover else first_f)
    if not continuous and pages[0].printed_number and pages[0].printed_number.isdigit():
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), pages[0].printed_number)
        sec._sectPr.append(pg)
