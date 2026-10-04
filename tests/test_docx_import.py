"""DOCX → EPUB: estrutura lida do Word."""

import re
import zipfile

from docx import Document

from bookbuilder.docx_import import docx_to_epub, read_docx
from bookbuilder.export.docx_writer import _hyperlink
from bookbuilder.model import Role, Run


def _make(path):
    d = Document()
    d.core_properties.title = "Livro do Word"
    d.add_heading("Capítulo Um", level=1)
    p = d.add_paragraph("Texto com ")
    p.add_run("negrito").bold = True
    p.add_run(" e um ")
    _hyperlink(p, "site do tribunal", Run("x"), "https://www.stf.jus.br")
    p.add_run(". Veja http://uqr.to/1yohc.")
    d.add_heading("Seção", level=2)
    d.add_paragraph("primeiro", style="List Number")
    d.add_paragraph("segundo", style="List Number")
    d.add_paragraph("item solto", style="List Bullet")
    d.add_paragraph("Uma citação.", style="Quote")
    d.add_heading("Capítulo Dois", level=1)
    d.add_paragraph("Fim.")
    d.save(path)


def test_le_estrutura(tmp_path):
    src = tmp_path / "a.docx"
    _make(src)
    res = read_docx(str(src), tmp_path / "fig")
    blocks = [b for p in res.project.pages for b in p.ordered_blocks()]
    assert len(res.project.pages) == 2  # um por capítulo
    assert [(b.role, b.heading_level) for b in blocks if b.role == Role.TITLE] == [
        (Role.TITLE, 1), (Role.TITLE, 2), (Role.TITLE, 1)]
    items = ["".join(r.text for r in b.runs) for b in blocks if b.role == Role.LIST_ITEM]
    assert items == ["1. primeiro", "2. segundo", "• item solto"]
    body = blocks[1]
    assert any(r.bold and r.text == "negrito" for r in body.runs)
    assert any(r.link == "https://www.stf.jus.br" and r.text == "site do tribunal" for r in body.runs)
    assert any(b.role == Role.QUOTE for b in blocks)
    assert res.project.export.title == "Livro do Word"


def test_gera_epub(tmp_path):
    src = tmp_path / "a.docx"
    _make(src)
    out = tmp_path / "a.epub"
    imp, er = docx_to_epub(str(src), str(out))
    assert er.chapters >= 2 and er.toc_entries == 3
    with zipfile.ZipFile(out) as z:
        html = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xhtml"))
    assert '<a href="https://www.stf.jus.br">site do tribunal</a>' in html
    assert '<a href="http://uqr.to/1yohc">' in html
    assert re.search(r"<blockquote><p>(<em>)?Uma citação\.(</em>)?</p></blockquote>", html)
